from pathlib import Path
from math import ceil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "engineering" / "hardware" / "alimentacao" / "assets"
OUT = ROOT / "docs" / "engineering" / "hardware" / "alimentacao" / "Referencia_tecnica_alimentacao_modular_VRLA_KiCad.docx"
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F3F8FC"
GRAY = "666666"
RED = "C00000"
GREEN = "2E7D32"
BLACK = "000000"


def font(size, bold=False):
    paths = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for p in paths:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


F14 = font(30, True)
F12 = font(24, True)
F10 = font(20, False)
F9 = font(17, False)
F8 = font(15, False)


def pc(color):
    return "#" + color if isinstance(color, str) and len(color) == 6 and not color.startswith("#") else color


def arrow(draw, p1, p2, color=NAVY, width=5):
    color = pc(color)
    draw.line([p1, p2], fill=color, width=width)
    x1, y1 = p1; x2, y2 = p2
    if abs(x2-x1) >= abs(y2-y1):
        s = 1 if x2 > x1 else -1
        draw.polygon([(x2, y2), (x2-18*s, y2-10), (x2-18*s, y2+10)], fill=color)
    else:
        s = 1 if y2 > y1 else -1
        draw.polygon([(x2, y2), (x2-10, y2-18*s), (x2+10, y2-18*s)], fill=color)


def center_text(draw, xy, text, ft=F9, fill=BLACK, spacing=5):
    box = draw.multiline_textbbox((0, 0), text, font=ft, spacing=spacing, align="center")
    x = xy[0] - (box[2]-box[0])/2
    y = xy[1] - (box[3]-box[1])/2
    draw.multiline_text((x, y), text, font=ft, fill=pc(fill), spacing=spacing, align="center")


def box(draw, rect, title, body="", fill=BLUE, outline=NAVY, body_fill=BLACK):
    draw.rounded_rectangle(rect, radius=18, fill=pc(fill), outline=pc(outline), width=4)
    x1, y1, x2, y2 = rect
    center_text(draw, ((x1+x2)/2, y1+42), title, F10, outline)
    if body:
        center_text(draw, ((x1+x2)/2, (y1+y2)/2+30), body, F8, body_fill)


def save_architecture():
    im = Image.new("RGB", (2200, 1180), "white"); d = ImageDraw.Draw(im)
    center_text(d, (1100, 58), "ARQUITETURA MODULAR - LIMITES DAS DUAS PCBs", F14, NAVY)
    box(d, (80, 170, 410, 390), "PAINEL SOLAR", "12 V nominal\nVmp tip. ~18 V\nVoc deve ser validado", "FFF2CC")
    box(d, (80, 620, 410, 840), "FONTE AUXILIAR", "Somente fonte DC\nisolada e certificada\n18 V tip. para carga", "FFF2CC")
    box(d, (560, 115, 1480, 920), "PLACA DE ALIMENTACAO\n(no quadro, perto da bateria)", "J1/J2 + fusivel + reversao + TVS + OR\nBQ24650: carregador solar CC/CV\nVRLA AGM 12 V externa\nTPS3700 + PMOS: LVD\nJ4: barramento PWR_BUS", "DDEBF7")
    box(d, (790, 420, 1240, 660), "BATERIA VRLA", "12 V AGM/GEL\nFusivel no polo +\nfora da PCB", "E2F0D9", GREEN)
    box(d, (1660, 270, 2110, 770), "PCB PRINCIPAL", "J_PWR / P4 selecionado\nprotecoes locais\nTPS62163: 5 V\nLDO/buck: 3V3\nSTM32, sensores, radio", "FCE4D6")
    arrow(d, (410, 280), (555, 280)); arrow(d, (410, 730), (555, 730))
    arrow(d, (1015, 385), (1015, 415), GREEN); arrow(d, (1240, 540), (1475, 540), GREEN); arrow(d, (1480, 540), (1655, 540), GREEN)
    d.rounded_rectangle((525, 80, 1515, 960), radius=26, outline=pc(NAVY), width=3)
    d.rounded_rectangle((1625, 235, 2145, 805), radius=26, outline=pc(NAVY), width=3)
    center_text(d, (1100, 1035), "A placa de alimentacao entrega tensao de bateria protegida, nominal 12 V, aproximadamente 11,3 a 14,7 V. Ela nao entrega 12,0 V regulados.", F9, RED)
    path = ASSETS / "01_arquitetura.png"; im.save(path); return path


def save_input_or():
    im = Image.new("RGB", (2200, 1180), "white"); d = ImageDraw.Draw(im)
    center_text(d, (1100, 55), "ENTRADAS PROTEGIDAS E OR DE FONTES - PLACA DE ALIMENTACAO", F14, NAVY)
    y1, y2 = 290, 760
    box(d, (70, 155, 300, 350), "J1 PV", "Painel solar", "FFF2CC")
    box(d, (70, 695, 300, 890), "J2 AUX", "Fonte DC 18 V\nde contingencia", "FFF2CC")
    box(d, (410, 155, 650, 350), "F1 / Q1", "Fusivel\nreversao", "DDEBF7")
    box(d, (410, 695, 650, 890), "F2 / Q2", "Fusivel\nreversao", "DDEBF7")
    box(d, (770, 155, 1010, 350), "TVS1 / D1", "surto +\nanti-retorno", "FCE4D6")
    box(d, (770, 695, 1010, 890), "TVS2 / D2", "surto +\nanti-retorno", "FCE4D6")
    box(d, (1190, 370, 1520, 655), "VIN_CHG", "nó comum apos OR\n5 a 28 V para BQ24650\nverificar Voc frio", "E2F0D9", GREEN)
    box(d, (1700, 370, 2115, 655), "U1 BQ24650", "carregador buck\nMPPT por tensao\nCC/CV para VRLA", "DDEBF7")
    for a,b in [((300,250),(405,250)),((650,250),(765,250)),((1010,250),(1165,430)),((300,790),(405,790)),((650,790),(765,790)),((1010,790),(1165,595)),((1520,512),(1695,512))]: arrow(d,a,b)
    center_text(d, (1100, 1000), "D1 e D2 podem ser Schottky no prototipo (perda termica) ou ideal-diode com MOSFET/controlador (menor perda). O OR bloqueia retroalimentacao; ele nao cria prioridade solar.", F9, RED)
    path = ASSETS / "02_entradas_or.png"; im.save(path); return path


def save_charger():
    im = Image.new("RGB", (2200, 1180), "white"); d = ImageDraw.Draw(im)
    center_text(d, (1100, 55), "CARREGADOR BQ24650 PARA VRLA - DIAGRAMA FUNCIONAL A IMPLEMENTAR", F14, NAVY)
    box(d, (60, 430, 330, 650), "VIN_CHG", "entrada apos OR\nCIN ceramicos", "FFF2CC")
    box(d, (460, 270, 830, 810), "U1 BQ24650", "VCC / REGN / BTST\nHIDRV / LODRV / PH\nSRP / SRN / VFB\nMPPSET / TS\nSTAT1 / STAT2\nTERM_EN", "DDEBF7")
    box(d, (970, 165, 1275, 380), "QH + QL", "NMOS sincronos\nloop de comutacao curto", "FCE4D6")
    box(d, (970, 700, 1275, 915), "R_SRP + L1", "shunt de corrente\nindutor de carga", "FCE4D6")
    box(d, (1430, 430, 1780, 650), "BAT_FUSED", "bateria VRLA\nterminal + com fusivel", "E2F0D9", GREEN)
    box(d, (930, 970, 1300, 1130), "REDE VFB", "resistores + NTC\nperfil SLA [R2]", "FFF2CC")
    box(d, (1480, 970, 1900, 1130), "STAT2 + Q3", "troca Vbulk/Vrecharge\nconforme circuito TI [R2]", "FFF2CC")
    arrow(d,(330,540),(455,540)); arrow(d,(830,420),(965,300)); arrow(d,(1120,380),(1120,695)); arrow(d,(1275,807),(1425,585)); arrow(d,(1120,915),(1120,965)); arrow(d,(1580,650),(1580,965))
    center_text(d,(1100,900),"Copiar a topologia e o layout do Typical Application do datasheet [R1]. O desenho mostra os blocos; nao substitui os pinos, desacoplamentos e valores do circuito oficial.",F9,RED)
    path=ASSETS/"03_bq24650.png"; im.save(path); return path


def save_lvd():
    im=Image.new("RGB",(2200,1180),"white"); d=ImageDraw.Draw(im)
    center_text(d,(1100,55),"LVD - CORTE AUTONOMO POR SUBTENSÃO DA BATERIA",F14,NAVY)
    box(d,(70,420,350,650),"BAT_FUSED", "da bateria\napos F3", "E2F0D9",GREEN)
    box(d,(510,315,870,750),"U2 TPS3700", "VDD em BAT_FUSED\nINA+ recebe divisor\nOUTA: dreno aberto\nOUTB: diagnostico OV", "DDEBF7")
    box(d,(510,840,870,1065),"DIVISOR R_U/R_D", "V_trip = 0,4 x\n(1 + R_U/R_D)\nbase sem histerese extra", "FFF2CC")
    box(d,(1050,315,1400,750),"QDRV + R_B", "transistor de gate\nOUTA baixo = desliga\nPMOS por padrao", "FCE4D6")
    box(d,(1570,315,1920,750),"Q_LVD PMOS", "S: BAT_FUSED\nD: PWR_BUS\nRGS: gate-source\nZDGS: protege VGS", "FCE4D6")
    box(d,(1970,420,2160,650),"J4 PWR_OUT", "PWR_BUS + GND\npara PCB principal", "DDEBF7")
    arrow(d,(350,535),(505,535)); arrow(d,(870,535),(1045,535)); arrow(d,(1400,535),(1565,535)); arrow(d,(1920,535),(1965,535))
    arrow(d,(690,750),(690,835));
    center_text(d,(1100,920),"O LVD continua funcionando sem MCU. Dimensionar dois limiares: corte (evita descarga profunda) e religamento (evita liga/desliga repetido); validar em carga e na temperatura de uso.",F9,RED)
    path=ASSETS/"04_lvd.png"; im.save(path); return path


def save_main():
    im=Image.new("RGB",(2200,1180),"white"); d=ImageDraw.Draw(im)
    center_text(d,(1100,55),"PCB PRINCIPAL - ENTRADA 12 V NOMINAL E RAILS DE 5 V / 3V3",F14,NAVY)
    box(d,(60,235,350,475),"J_PWR", "da placa de alimentacao\nPWR_BUS/GND", "DDEBF7")
    box(d,(60,710,350,950),"J_DC P4", "fonte 12 V externa\nversao simples", "FFF2CC")
    box(d,(500,380,880,805),"SELECAO DE ENTRADA", "A: jumper 0R\nB: OR por diodos\nC: variantes exclusivas\nNunca curto direto P4 + J_PWR", "FCE4D6")
    box(d,(1030,380,1400,805),"U3 TPS62163", "buck fixo 5 V / 1 A\nVIN 3 a 17 V\nL, CIN, COUT, EN, PG", "DDEBF7")
    box(d,(1560,220,2050,520),"RAIL 5V", "radio 5 V, sensores\nUSB se aplicavel\nmedir pico de corrente", "E2F0D9",GREEN)
    box(d,(1560,690,2050,990),"U4 3V3", "MCP1700 se I3V3 baixo\nou buck se carga alta\nSTM32 + sensores 3V3", "DDEBF7")
    arrow(d,(350,355),(495,505)); arrow(d,(350,830),(495,675)); arrow(d,(880,595),(1025,595)); arrow(d,(1400,595),(1555,370)); arrow(d,(1805,520),(1805,685))
    center_text(d,(1100,1090),"A variacao natural da VRLA (aprox. 11,3 a 14,7 V enquanto o LVD permite) esta dentro de VIN do TPS62163. Nao adicionar regulador de 12 V a menos que exista uma carga que realmente exija 12,0 V.",F9,RED)
    path=ASSETS/"05_pcb_principal.png"; im.save(path); return path


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def borders(cell, color="B7C9D6", size="6"):
    tcPr = cell._tc.get_or_add_tcPr(); tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = tcBorders.find(qn(tag))
        if el is None: el = OxmlElement(tag); tcBorders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color)


def set_cell(cell, text, bold=False, color=BLACK, size=8.6, fill=None):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r=p.add_run(text); r.bold=bold; r.font.name="Arial"; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill: shade(cell,fill)
    borders(cell)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),"true"); trPr.append(tblHeader)


def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    hdr=t.rows[0]
    for c,h in zip(hdr.cells,headers): set_cell(c,h,True,"FFFFFF",8.4,NAVY)
    set_repeat_table_header(hdr)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for c,v in zip(cells,row): set_cell(c,str(v),False,BLACK,8.3, PALE if idx%2==0 else "FFFFFF")
    if widths:
        for row in t.rows:
            for cell,w in zip(row.cells,widths): cell.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t


def set_keep(paragraph):
    pPr=paragraph._p.get_or_add_pPr(); el=OxmlElement('w:keepNext'); pPr.append(el)


def add_heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(11 if level==1 else 7); p.paragraph_format.space_after=Pt(5); set_keep(p)
    r=p.add_run(text); r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); r.font.color.rgb=RGBColor(0,0,0)
    return p


def add_p(doc, text="", bold_start=None, style=None, align=None):
    p=doc.add_paragraph(style=style) if style else doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.08
    if bold_start and text.startswith(bold_start):
        r=p.add_run(bold_start); r.bold=True; r.font.name='Arial'; r.font.size=Pt(9.6)
        r=p.add_run(text[len(bold_start):]); r.font.name='Arial'; r.font.size=Pt(9.6)
    else:
        r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(9.6)
    if align: p.alignment=align
    return p


def bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(0); p.paragraph_format.left_indent=Inches(0.18+0.18*level)
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(9.4); return p


def note(doc, title, text, kind="info"):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; cell=t.cell(0,0)
    fill={"info":"EAF3F8","warn":"FFF2CC","risk":"FCE4D6"}[kind]; shade(cell,fill); borders(cell, "9FBAD0")
    p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(title+". "); r.bold=True; r.font.name='Arial'; r.font.size=Pt(9.2); r.font.color.rgb=RGBColor.from_string(NAVY if kind=='info' else RED if kind=='risk' else "7F6000")
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.italic=True; r.font.name='Arial'; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(GRAY)


def add_fig(doc, path, caption_text, width=6.7):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    p.add_run().add_picture(str(path), width=Inches(width)); caption(doc, caption_text)


def page_number(section):
    footer=section.footer; p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run("Pagina "); r.font.name='Arial'; r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)


def build():
    imgs=[save_architecture(),save_input_or(),save_charger(),save_lvd(),save_main()]
    d=Document(); sec=d.sections[0]
    sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.66); sec.right_margin=Inches(.66)
    page_number(sec)
    styles=d.styles
    styles['Normal'].font.name='Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); styles['Normal'].font.size=Pt(9.6)
    for n,sz in [('Title',25),('Heading 1',15),('Heading 2',12),('Heading 3',10.5)]:
        st=styles[n]; st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); st.font.size=Pt(sz); st.font.color.rgb=RGBColor(0,0,0); st.font.bold=True

    # Cover
    p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(78); p.paragraph_format.space_after=Pt(12)
    r=p.add_run('Referencia tecnica\nde alimentacao modular com VRLA AGM'); r.font.name='Arial'; r.font.color.rgb=RGBColor(0,0,0)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(24)
    r=p.add_run('Guia sequencial para esquematico, dimensionamento e implementacao no KiCad'); r.font.name='Arial'; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(NAVY)
    note(d,'Escopo', 'Estacao remota de baixo consumo, alimentada prioritariamente por painel solar e bateria VRLA de 12 V. A placa de alimentacao fica separada da PCB principal e entrega um barramento de bateria protegido. A rede eletrica entra somente por uma fonte DC externa, isolada e certificada, como contingencia.', 'info')
    add_p(d,'Versao 1.0 - 18 de setembro de 2026',align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(d,'Este documento e uma referencia de engenharia. Antes de fabricar, fechar painel, bateria, perfil de carga, corrente de pico, temperatura e footprints com os respectivos datasheets.',align=WD_ALIGN_PARAGRAPH.CENTER)
    d.add_page_break()

    add_heading(d,'1. Decisao de arquitetura e limites do sistema')
    add_p(d,'A arquitetura recomendada e modular: a placa de alimentacao e instalada no quadro, perto da bateria, painel e conectores de maior corrente. A PCB principal abriga apenas a conversao local, MCU, sensores e comunicacao. Essa separacao reduz ruido de comutacao perto dos sensores, facilita manutencao e permite usar a mesma PCB principal tanto com a placa solar quanto com uma fonte DC 12 V externa.')
    add_fig(d,imgs[0],'Figura 1 - Visao geral. Setas verdes representam caminho de energia; os limites das PCBs evitam que corrente de carga atravesse conectores pequenos.')
    add_table(d,['Bloco','Onde fica','Funcao','Observacao de projeto'],[
        ['J1, J2, F1-F3, TVS, OR','Placa de alimentacao','Protegem fontes, isolam retornos e limitam falhas','Usar conectores e trilhas para corrente de carga.'],
        ['U1 BQ24650, MOSFETs e indutor','Placa de alimentacao','Carrega VRLA usando painel / fonte auxiliar','Nao e modulo pronto; e conversor buck controlado por CI.'],
        ['Bateria VRLA e fusivel F3','Externo','Reserva de energia noturna','F3 deve ficar fisicamente proximo ao polo positivo.'],
        ['U2 TPS3700 + PMOS','Placa de alimentacao','LVD: corta carga em subtensao','Opera mesmo que MCU e PCB principal estejam desligados.'],
        ['U3 e U4','PCB principal','Convertem PWR_BUS para 5 V e 3V3','Entrada tolera faixa da bateria, nao exige 12,0 V fixos.'],
    ],[1.15,1.35,2.15,2.0])
    note(d,'Decisao importante','PWR_BUS e um barramento nominal de 12 V, e nao uma saida regulada de 12,0 V. Na operacao normal ele acompanha aproximadamente 11,3 a 14,7 V, conforme o LVD, a carga, a fase de carga e a temperatura.', 'warn')

    add_heading(d,'2. Dados que precisam ser fechados antes dos valores finais')
    add_p(d,'O esquema pode ser organizado agora, mas os valores de fusivel, shunt, indutor, bateria e painel nao devem ser escolhidos por intuicao. Preencha a tabela apos medir a primeira PCB, inclusive os picos do radio e do cartao SD.')
    add_table(d,['Parametro','Como obter','Uso direto no projeto','Valor inicial'],[
        ['P_media da estacao','Medir V e I em ciclo completo (sleep, leitura, TX)','Capacidade da bateria e painel','A medir'],
        ['I_pico 5 V e 3V3','Osciloscopio/medidor durante TX e escrita','U3/U4, capacitores, conectores','A medir'],
        ['Horas sem sol (H_noite)','Pior noite local + margem','Energia noturna','12 a 14 h para exemplo'],
        ['HSP local','Base solar da cidade / pior mes','Potencia do painel','A levantar'],
        ['Perfil VRLA','Datasheet exato da bateria e temperatura','Vbulk, Vfloat/Vrecharge, Icharge','Obrigatorio'],
        ['Voc do painel frio','Datasheet do painel + menor temperatura','Limite 28 V do BQ24650','Obrigatorio'],
    ],[1.35,2.0,2.1,1.2])
    note(d,'Seguranca','Painel e bateria devem ficar fora do sol direto e de fontes de calor. A fonte auxiliar deve ser uma fonte AC/DC comercial certificada: nenhuma tensao de rede entra nesta PCB. VRLA e selada, mas pode aquecer e liberar gases em condicoes de falha/sobrecarga; usar caixa ventilada, fusivel e perfil de carga correto.', 'risk')

    add_heading(d,'3. Dimensionamento de energia: bateria e painel')
    add_heading(d,'3.1 Procedimento de medicao de consumo',2)
    bullet(d,'Meça corrente em cada estado: Stop/Standby do STM32, sensores ligados, conversao, radio transmitindo, radio recebendo, SD gravando e LEDs.')
    bullet(d,'Para um ciclo de duracao T, calcule P_media = soma(V x I_i x t_i) / T. Use a tensao real de cada rail e considere eficiencia do buck.')
    bullet(d,'Repita a medida na tensao de bateria baixa e com a temperatura esperada. O pico determina margens de corrente; a media determina autonomia.')
    add_heading(d,'3.2 Capacidade da bateria',2)
    add_p(d,'Energia noturna: E_noite [Wh] = P_media [W] x H_noite [h]. Capacidade nominal minima: C_Ah >= E_noite / (V_nom x DoD_uso x eta_caminho). Para preservar VRLA, use DoD_uso inicial de 0,50 e eficiencia de caminho de 0,85. Multiplique por margem de envelhecimento, temperatura e dias nublados de 1,25 a 1,50.')
    add_table(d,['Exemplo didatico','Calculo','Resultado'],[
        ['P_media = 0,50 W; H_noite = 14 h','E_noite = 0,50 x 14','7 Wh'],
        ['V_nom = 12 V; DoD = 0,50; eta = 0,85','C_min = 7 / (12 x 0,50 x 0,85)','1,37 Ah'],
        ['Margem 1,5','1,37 x 1,5','2,1 Ah; selecionar acima e validar descarga'],
        ['P_media = 2,0 W; H_noite = 14 h','28 / (12 x 0,50 x 0,85)','5,49 Ah; tipicamente 9-12 Ah apos margem'],
    ],[2.0,3.0,1.65])
    add_heading(d,'3.3 Potencia do painel',2)
    add_p(d,'Com E_dia como energia que deve ser reposta por dia, estime P_painel >= E_dia / (HSP x eta_solar). Comece com eta_solar = 0,65 para perdas de temperatura, poeira, cabo, conversao e carga. A corrente disponivel do painel tambem precisa respeitar a corrente programada do carregador.')
    note(d,'Exemplo','Se a estacao consome 7 Wh na noite e 7 Wh no dia, E_dia e pelo menos 14 Wh. Com HSP = 3,5 h e eta_solar = 0,65, P_painel >= 6,2 W. Escolha um painel maior apos considerar dias nublados e a recarga da bateria.', 'info')

    d.add_page_break()
    add_heading(d,'4. Etapa 1 - Entradas protegidas e OR de fontes')
    add_p(d,'A placa de alimentacao possui duas entradas de energia de carga: J1 para painel solar e J2 para uma fonte DC auxiliar, usada apenas quando a energia solar nao basta. Cada caminho e protegido independentemente antes de se encontrar em VIN_CHG. O objetivo do OR e impedir que uma fonte alimente a outra, por exemplo, que a bateria descarregue para o painel escuro por meio do carregador.')
    add_fig(d,imgs[1],'Figura 2 - Blocos de entrada. O caminho fisico de cada fonte e: conector -> fusivel -> protecao de polaridade -> TVS -> bloqueio de retorno -> VIN_CHG.')
    add_heading(d,'4.1 Circuito de cada entrada',2)
    add_table(d,['Ref.','Ligacao no KiCad','Funcao','Dimensionamento / observacao'],[
        ['J1 / J2','Conector de 2 vias: + e GND','Entrada de painel / fonte auxiliar','Corrente nominal do conector >= 2 x corrente maxima prevista.'],
        ['F1 / F2','Em serie no positivo, logo apos conector','Limita energia em curto do cabo/placa','I_fusivel >= 1,25 x I_continua e abaixo da capacidade do cabo/conector. Avaliar inrush.'],
        ['Q1 / Q2','PMOS de polaridade reversa, gate-source com resistores','Evita dano se + e - forem invertidos','VDS >= margem sobre Voc/aux; VGS protegido se necessario; RDS(on) avaliado em VGS real.'],
        ['TVS1 / TVS2','Do positivo protegido para GND, apos fusivel','Prende surtos de cabo','VRWM acima da tensao normal maxima; Vclamp abaixo do limite dos CIs. Ver curva de pulso.'],
        ['D1 / D2','Em serie apos cada caminho, anodo na fonte e catodo em VIN_CHG','OR e anti-retorno','Schottky e simples; ideal-diode reduz queda e calor.'],
        ['CIN','VIN_CHG para GND, muito perto de QH/QL do carregador','Fecha loop pulsante do buck','Usar ceramicos X7R; seguir valores/ESR do datasheet [R1].'],
    ],[.55,2.25,1.6,2.25])
    add_heading(d,'4.2 Tres alternativas para o OR',2)
    add_table(d,['Alternativa','Como implementar','Vantagem','Desvantagem'],[
        ['A. Schottky','D1 e D2 Schottky em serie para VIN_CHG','Menor custo e facil de depurar','Queda V_F e perda P = I x V_F; pode aquecer.'],
        ['B. PMOS ideal diode','PMOS orientado para bloquear retorno, circuito de gate adequado ou controlador','Queda aproximada I x RDS(on), mais eficiente','Mais detalhes de gate, validacao de transitorios.'],
        ['C. Controlador ideal-diode','CI controlador + MOSFET por entrada','Melhor para corrente maior e hot-swap','Maior custo e area; seguir referencia do fabricante [R6].'],
    ],[1.2,2.2,1.5,1.75])
    note(d,'O que OR nao faz','OR permite que a fonte com maior tensao efetiva contribua primeiro e bloqueia retorno. Ele nao garante “prioridade solar”. Para uma prioridade formal seria necessario controle adicional. Para o TCC, deixar a fonte auxiliar desconectada no uso normal e o caminho mais simples e verificavel.', 'warn')
    add_heading(d,'4.3 Exemplo de perdas e escolha inicial',2)
    add_p(d,'Se I_charge = 0,8 A e D1 for uma Schottky de 0,45 V no ponto de operacao, P_D = 0,8 x 0,45 = 0,36 W. Isso ja exige avaliar encapsulamento e temperatura. Com PMOS de 50 mOhm, P_Q = I^2 x R = 0,8^2 x 0,05 = 0,032 W, mas o circuito fica mais sensivel a layout e tensao gate-source.')

    add_heading(d,'5. Etapa 2 - Carregador solar BQ24650 e perfil VRLA')
    add_p(d,'O BQ24650RVAR e um controlador buck sincrono de carga, nao um modulo fechado. Ele comanda dois NMOS externos, indutor, resistor shunt e rede de realimentacao. A entrada opera de 5 V a 28 V e o CI implementa limitacao por tensao de entrada para operar painel solar proximo ao ponto programado. O CI suporta quimica chumbo-acida, mas a rede de VFB e STAT2 deve seguir a aplicacao SLA oficial, nao o circuito Li-ion generico.')
    add_fig(d,imgs[2],'Figura 3 - Blocos do carregador. Reproduzir a topologia, o desacoplamento e o layout do “Typical Application” [R1] e a Figura 2 da nota de aplicacao SLA [R2].')
    add_heading(d,'5.1 O que cada grupo de pinos exige no esquema',2)
    add_table(d,['Grupo','Conexao / circuito externo','Papel na operacao'],[
        ['VCC, REGN, GND','VIN_CHG, capacitores locais de REGN/VCC conforme [R1], plano GND','Alimentacao e driver interno. Desacoplamento deve ficar no proprio CI.'],
        ['HIDRV, LODRV, PH, BTST','Dois NMOS externos, bootstrap e loop de comutacao curto','Formam o buck sincrono que regula a corrente de carga.'],
        ['SRP, SRN','Nas extremidades Kelvin de RSR','Mede queda no shunt e regula I_charge. Nao compartilhar essa medicao com caminho de potencia.'],
        ['VFB','Rede de resistores, NTC e chave STAT2/Q3 da referencia SLA','Define Vbulk e Vrecharge / flutuacao conforme perfil VRLA.'],
        ['MPPSET','Divisor resistivo de VIN_CHG e, opcionalmente, transistor de disable','Programa a tensao de regulacao da entrada, o MPPT por tensao.'],
        ['TS','NTC na bateria ou rede fixa somente se o perfil permitir','Qualificacao termica. Em ambiente externo, preferir NTC fisicamente perto da VRLA.'],
        ['STAT1, STAT2','Saidas open-drain com pull-up adequado; LEDs e/ou sinais de status','Indicam estado. STAT2 tambem participa da rede SLA de VFB.'],
        ['TERM_EN','Nivel conforme datasheet e algoritmo SLA','Habilita terminacao conforme a estrategia adotada.'],
    ],[1.25,3.05,2.25])
    add_heading(d,'5.2 Corrente de carga, shunt e indutor',2)
    add_p(d,'No BQ24650, a referencia de corrente e aproximadamente 40 mV no shunt. Portanto RSR = 0,040 / I_charge. A corrente continua deve respeitar o fabricante da VRLA; a nota [R2] cita como orientacao 0,1 C a 0,3 C no estagio de corrente constante, mas o datasheet da bateria escolhida prevalece.')
    add_table(d,['Exemplo','Calculo','Resultado / observacao'],[
        ['VRLA 7 Ah, I_charge = 0,2 C','I_charge = 7 x 0,2','1,4 A'],
        ['RSR para 1,4 A','0,040 / 1,4','28,6 mOhm; usar valor padrao e recalcular'],
        ['Perda no shunt','P = I^2 x R','~0,056 W nominal; especificar margem e Kelvin'],
        ['Indutor','Usar procedimento da secao de selecao de L em [R1]','Isat acima do pico; DCR baixo; nucleo nao saturar.'],
        ['MOSFETs QH/QL','VDS e corrente com margem, RDS(on), Qg e termica','Nao escolher apenas por corrente de catalogo.'],
    ],[1.6,2.2,2.75])
    add_heading(d,'5.3 Tensao de entrada / MPPT',2)
    add_p(d,'O BQ24650 regula a entrada por tensao, nao faz MPPT digital completo. Para um painel de 12 V nominal com Vmp tipico perto de 18 V, comece com MPPSET perto da Vmp do painel, mas feche pelo datasheet e por ensaio em irradiancia/temperatura. A formula do divisor vem do datasheet [R1], que usa referencia interna de 1,2 V: V_MPP ~= 1,2 x (1 + Rtop/Rbot). A tensao de circuito aberto do painel no frio precisa permanecer dentro do limite de entrada do CI, inclusive transitorios.')
    add_heading(d,'5.4 Perfil de carga VRLA - ponto critico',2)
    add_p(d,'A nota TI SLVA437A [R2] adapta o BQ24650 para SLA selada: corrente constante ate Vbulk e depois tensao constante. A logica baseada em STAT2 e Q3 muda a rede de realimentacao entre Vbulk e Vrecharge. No exemplo da TI para bateria de 2,4 Ah, Vbulk = 14,85 V e Vrecharge = 14,1 V a 25 C; estes numeros sao apenas exemplo, nao valores universais para a sua bateria.')
    note(d,'Obrigatorio antes da fabricacao','Escolher o modelo exato de VRLA e copiar do seu datasheet: tensao de uso ciclico/absorpcao, tensao de flutuacao, coeficiente com temperatura, corrente maxima de carga e faixa de temperatura. Usar valores errados encurta a vida e pode provocar aquecimento.', 'risk')

    d.add_page_break()
    add_heading(d,'6. Etapa 3 - Bateria, fusivel e LVD autonomo')
    add_p(d,'O LVD (Low Voltage Disconnect) desliga a carga da estacao quando a bateria cai abaixo de um limiar seguro. Ele nao regula tensao e nao substitui um carregador: seu papel e preservar a VRLA contra descarga profunda. O TPS3700 fica alimentado antes do PMOS de corte, portanto pode monitorar e religar o sistema mesmo com a PCB principal desligada.')
    add_fig(d,imgs[3],'Figura 4 - LVD conceitual. OUTA e open-drain; a polaridade do driver deve garantir PMOS desligado se U2 perder alimentacao ou o cabo da PCB principal estiver ausente.')
    add_heading(d,'6.1 Caminho e componentes',2)
    add_table(d,['Ref.','Conexao','Por que existe','Criterio de escolha'],[
        ['F3','Polo + da bateria, fisicamente perto do terminal','Protege cabo e placa contra curto de alta energia','Fusivel de carro/inline ou porta-fusivel adequado; corrente e cabo dimensionados juntos.'],
        ['U2 TPS3700DDCT','VDD em BAT_FUSED; GND comum','Supervisor de janela, 1,8-18 V, Iq tip. 5,5 uA [R3]','SOT-23-6 mais soldavel que WSON.'],
        ['R_U / R_D','BAT_FUSED -> INA+ -> GND','Escala tensao da bateria para referencia de 0,4 V','1% e valores que limitam corrente de divisor; verificar tolerancia.'],
        ['QDRV + R_B','OUTA controla driver do gate','Converte open-drain em comando robusto de PMOS','Escolher polaridade fail-safe: sem drive, Q_LVD OFF.'],
        ['Q_LVD','Source em BAT_FUSED; drain em PWR_BUS','Chave high-side de corte da carga','VDS >= margem da bateria/transitorios; RDS(on) por perda I^2R.'],
        ['RGS / ZDGS','Gate-source de Q_LVD','RGS desliga por padrao; zener limita |VGS|','ZDGS abaixo do max. |VGS| do PMOS e acima do drive necessario.'],
        ['F4 / PTC opcional','Depois de Q_LVD, antes de J4','Protege cabo interplacas','Selecionar contra falha do cabo, sem disparar nos picos normais.'],
    ],[.8,2.1,1.9,2.0])
    add_heading(d,'6.2 Calculo do limiar e histerese',2)
    add_p(d,'Sem rede externa de histerese, o limiar de subida aproximado e V_trip = 0,400 x (1 + R_U/R_D). Para um limiar nominal de 11,5 V, a razao R_U/R_D e 11,5/0,4 - 1 = 27,75. Exemplo inicial: R_D = 10 kOhm e R_U = 280 kOhm produzem aproximadamente 11,6 V. O TPS3700 possui histerese interna tipica de 5,5 mV na entrada, que vira apenas cerca de 0,16 V no barramento para essa razao; isso pode ser insuficiente contra bounce de carga.')
    add_p(d,'Por isso, preveja footprint para R_HYS entre a saida/driver e o no do divisor, seguindo a topologia de supervisor do datasheet [R3], ou use uma rede dedicada de realimentacao. Feche R_HYS por simulacao e ensaio. Como ponto de partida de requisitos, defina V_cut = 11,3 a 11,5 V e V_reconnect = 12,0 a 12,4 V, mas aprove somente depois de descarregar a bateria real sob a corrente da estacao.')
    add_heading(d,'6.3 OUTA e OUTB',2)
    add_p(d,'Use OUTA para o limite inferior: ele vai a nivel baixo quando INA+ cai abaixo do limiar. OUTB pode ficar como diagnostico de sobretensao usando INB- e outro divisor, ou nao ser usado nesta versao. Nao dependa do STM32 para a decisao de corte; o MCU apenas pode receber um sinal BAT_LOW antes do LVD por um divisor/ADC separado.')
    note(d,'Limite de seguranca','Nao conectar uma bateria VRLA sem F3 nem permitir que toda a corrente de falha atravesse JST-SH. A bateria pode fornecer corrente alta mesmo “pequena”.', 'risk')

    add_heading(d,'7. Etapa 4 - Saida da placa de alimentacao e conector interplacas')
    add_p(d,'J4 entrega PWR_BUS e GND para a PCB principal. Para manter a modularidade, a placa de alimentacao precisa funcionar por conta propria: BQ24650, TPS3700 e o LVD nao exigem conexao com MCU. Um chicote de sinais e opcional e deve ficar separado do caminho de potencia.')
    add_table(d,['Pino J4 sugerido','Direcao','Nivel / funcao','Observacao'],[
        ['1 PWR_BUS','Saida','11,3-14,7 V nominal','Pino/cabo para corrente continua e pico com margem.'],
        ['2 GND','Retorno','0 V comum','Mesmo calibre do positivo.'],
        ['3 CHG_STAT (opcional)','Saida open-drain','Estado do carregador','Pull-up em 3V3 na PCB principal ou na power board, nunca exceder rating.'],
        ['4 BAT_LOW_PRE (opcional)','Saida digital','Aviso antes do LVD','Permite MCU salvar dados antes de corte.'],
        ['5-6 reserva','-','NTC/telemetria futura','Nao obrigatorios na primeira versao.'],
    ],[1.5,1.0,2.0,2.3])
    add_p(d,'Para PWR_BUS e GND, usar conector travado com especificacao de corrente apropriada (por exemplo familias JST-VH, Molex Micro-Fit ou borne de parafuso, conforme corrente). JST-SH e JST-GH sao adequados para sinais de baixa corrente; nao para alimentar a estacao se a corrente prevista exceder o rating do conector.')
    add_heading(d,'7.1 Queda no cabo',2)
    add_p(d,'Calcule V_drop = 2 x L_cabo x R_metro x I. O fator 2 contempla ida e retorno. Exemplo: 1,0 m de cabo 0,35 mm2, aproximadamente 0,05 Ohm/m por condutor, e 0,8 A resulta em 2 x 1 x 0,05 x 0,8 = 0,08 V. Avalie tambem a dissipacao do cabo P = I^2 x R_total.')
    note(d,'Terra e ruido','Use GND comum entre placas pelo par de potencia. Na PCB principal, separe fisicamente o loop de entrada/buck dos sinais analogicos; nao use fio fino de sinal como retorno de potencia.', 'info')

    add_heading(d,'8. Etapa 5 - PCB principal: selecao de entrada e conversao para 5 V')
    add_p(d,'A PCB principal deve aceitar tanto J_PWR (placa de alimentacao) quanto uma fonte DC 12 V via P4 para instalacoes simples. Essas duas entradas nunca devem ser unidas diretamente sem uma politica de selecao: uma fonte pode tentar alimentar a outra. O barramento da bateria ja varia; o conversor de 5 V foi escolhido para suportar a faixa esperada.')
    add_fig(d,imgs[4],'Figura 5 - PCB principal. As alternativas de selecao P4/J_PWR sao exclusivas no custo minimo ou OR se coexistencia simultanea for realmente necessaria.')
    add_heading(d,'8.1 Tres formas de conectar P4 e J_PWR',2)
    add_table(d,['Forma','Implementacao','Quando usar','Vantagem / desvantagem'],[
        ['A. Jumper 0R ou solder bridge','JP1 seleciona somente J_PWR ou somente P4 antes de VIN_MAIN','TCC inicial; uma fonte por vez','Mais barato e claro. Exige configurar antes de energizar.'],
        ['B. OR de entradas','Diodos Schottky ou ideal diode entre P4/J_PWR e VIN_MAIN','Hot-swap ou fontes simultaneas','Automatico, mas aumenta custo/perda e exige testes.'],
        ['C. Variantes de montagem','DNP em P4 ou J_PWR, conforme produto','Produtos separados','Minima complexidade e zero ambiguidade; nao e hot-swap.'],
    ],[1.35,2.2,1.55,1.8])
    add_heading(d,'8.2 Protecao local da entrada',2)
    add_p(d,'Mesmo recebendo a alimentacao protegida, inclua F5 ou PTC de baixa corrente, TVS adequado ao barramento, capacitor bulk e ceramico em VIN_MAIN. A protecao local protege o cabo, o conector e o loop de entrada do buck. Ela nao substitui F3 perto da bateria.')
    add_heading(d,'8.3 U3 TPS62163DSGR - 5 V / 1 A',2)
    add_p(d,'TPS62163DSGR e a versao fixa de 5 V, 1 A, em WSON 2 x 2 mm. Opera com VIN recomendado de 3 a 17 V, possui modo de baixo consumo e Iq tipico de 17 uA [R4], portanto aceita a faixa da VRLA. No esquema, colocar CIN entre VIN e GND, indutor entre SW e VOUT, COUT de VOUT a GND, EN definido (pull-up/pull-down conforme desejado), PG com pull-up na rail adequada e VOS ligado a VOUT. Para a versao fixa, FB vai a GND conforme datasheet.')
    add_table(d,['Elemento de U3','Como dimensionar','Nota de implementacao'],[
        ['Corrente','I_5V_cont e I_5V_pico medidos; manter margem abaixo de 1 A','Se radio/atuador exceder ou tiver grande pico, trocar por buck de maior corrente.'],
        ['Indutor L3','Valor e Isat conforme tabela do datasheet [R4]','Isat maior que limite/pico; DCR afeta eficiencia.'],
        ['CIN/COUT','Valores X7R e tensao conforme [R4]','Considerar derating de capacitancia DC; manter proximos aos pinos.'],
        ['Termica','P_loss aproximada e thetaJA do layout','WSON exige pad exposto e vias termicas como recomendado.'],
        ['Package','TPS62163DSGR','Nao confundir com “DGKR”: DSGR e WSON fixo 5 V; TPS62160DGKR e versao ajustavel VSSOP.'],
    ],[1.35,2.45,3.1])

    add_heading(d,'9. Etapa 6 - Rail de 3V3 e distribuicao de carga')
    add_p(d,'A partir do 5 V, gerar 3V3 para STM32 e sensores. Um LDO de baixo Iq reduz custo e ruido quando a corrente 3V3 e pequena; um segundo buck reduz perda quando a corrente ou a diferenca termica e maior. O regulador nao deve ser escolhido somente pelo “maximo de corrente” de catalogo: medir o pior pico do radio e validar a temperatura.')
    add_table(d,['Opcao','Circuito','Quando faz sentido','Calculo principal'],[
        ['MCP1700-3302','5V -> CIN -> LDO -> COUT -> 3V3','MCU/sensores com I3V3 baixo, por exemplo <100 mA medio','P_LDO = (5 - 3,3) x I3V3. A 100 mA: 0,17 W.'],
        ['Buck 3V3','5V -> buck sincrono -> L/C -> 3V3','Radio, SD ou pico 3V3 relevante','Perda menor; custo, layout e ripple maiores.'],
        ['Rail com load switch','3V3 -> chave -> sensor/radio','Cargas que podem dormir de verdade','I_sleep reduzido e controle pelo MCU.'],
    ],[1.4,2.0,2.0,1.5])
    add_p(d,'No MCP1700, seguir os capacitores e requisitos de estabilidade do datasheet [R5]. O calculo termico de primeira ordem e DeltaT ~= P_LDO x thetaJA, mas a validacao deve ser feita na PCB real. Separar capacitor local de cada IC, radio e sensor, e prever bulk proximo a cargas pulsantes.')
    add_heading(d,'9.1 Interface com STM32: o que precisa e o que nao precisa')
    add_table(d,['CI / bloco','Precisa de MCU?','Ligacoes opcionais e recomendadas'],[
        ['BQ24650','Nao. Opera autonomamente.','STAT1/STAT2 para diagnostico; MPPSET pode ser reduzido por transistor para disable planejado. Nao aplicar tensao de MCU fora do rating.'],
        ['TPS3700 / LVD','Nao. Deve proteger mesmo sem firmware.','BAT_LOW_PRE para ADC/interrupt e permitir salvar dados antes do corte.'],
        ['TPS62163','Nao para regular 5 V.','EN para desligar rail em modo economia; PG para diagnostico/boot sequenciado.'],
        ['U4 3V3','Nao se sempre ligada.','EN de buck/load switch para sensores com duty cycle.'],
    ],[1.45,1.55,3.9])
    note(d,'Nivel logico','STAT e PG sao tipicamente saidas open-drain. O pull-up pode ficar em 3V3 na PCB principal, desde que o fio de sinal e o CI suportem o arranjo. Para sinais que atravessam placas, adicionar resistor serie (por exemplo 100 Ohm como footprint) e GND de referencia.', 'warn')

    add_heading(d,'10. Modos de funcionamento do sistema')
    add_table(d,['Estado','Fluxo de energia','Acao dos blocos','Comportamento esperado'],[
        ['Dia, painel suficiente','Painel -> BQ -> bateria; bateria/PWR_BUS -> PCB','BQ regula entrada e carga; LVD permanece fechado','Estacao opera e bateria recarrega.'],
        ['Dia, painel fraco','Painel limitado; bateria complementa barramento','BQ reduz corrente pela regulacao de entrada','Sem puxar painel abaixo do MPPSET programado.'],
        ['Noite','VRLA -> F3 -> LVD -> PCB','BQ dorme/nao carrega; U3/U4 alimentam cargas','Autonomia dada pelo dimensionamento.'],
        ['Bateria baixa','VRLA atinge V_cut','TPS3700 desliga PMOS LVD','PCB perde alimentacao; preserva bateria.'],
        ['Sol retorna / recarga','BQ carrega VRLA acima V_reconnect','TPS3700 religa PMOS apos histerese','Evita ciclo rapido liga/desliga.'],
        ['Emergencia','Fonte AUX -> OR -> BQ -> VRLA','AUX fornece energia de carga','Nao ligar fonte DC 12 V comum esperando carga completa da VRLA.'],
    ],[1.3,2.0,2.0,1.6])
    add_p(d,'A fonte auxiliar deve ser escolhida para a tensao de entrada e corrente de carga do BQ24650. Uma fonte 12 V comum pode alimentar a PCB principal pela entrada P4 na versao simples, mas normalmente nao fornece a margem de tensao necessaria para carregar uma VRLA de 12 V ate o perfil correto.')

    add_heading(d,'11. Organizando o projeto no KiCad')
    add_heading(d,'11.1 Hierarquia recomendada',2)
    add_table(d,['Folha','Conteudo','Labels de rede principais'],[
        ['PWR_01_Input_OR','J1/J2, F1/F2, Q1/Q2, TVS, D1/D2','PV_RAW, AUX_RAW, VIN_CHG, GND'],
        ['PWR_02_BQ24650','U1, QH/QL, L1, RSR, VFB, MPPSET, TS','VIN_CHG, BAT_FUSED, PH, STAT1, STAT2'],
        ['PWR_03_Battery_LVD','J3 bateria, F3, U2, QDRV, Q_LVD, J4','BAT_FUSED, PWR_BUS, BAT_LOW_PRE, GND'],
        ['MAIN_01_Input','J_PWR, J_DC/P4, JP1/OR, F5, TVS, VIN_MAIN','PWR_BUS, VIN_MAIN, GND'],
        ['MAIN_02_5V','U3, L3, CIN, COUT, EN, PG','VIN_MAIN, 5V, PWR_GOOD_5V, GND'],
        ['MAIN_03_3V3_Loads','U4/buck, sensores, load switches','5V, 3V3, SENSOR_EN, GND'],
    ],[1.5,3.25,2.05])
    add_heading(d,'11.2 Regras de esquema',2)
    bullet(d,'Use labels globais somente para rails realmente comuns. PH/SW, SRP/SRN e VFB devem ficar locais e claros.')
    bullet(d,'Marque DNP (Do Not Populate) para alternativas: D1/D2 Schottky versus ideal diode, R_HYS, P4/J_PWR em variantes e pull-ups de sinais.')
    bullet(d,'Coloque test points em VIN_CHG, BAT_FUSED, PWR_BUS, VIN_MAIN, 5V, 3V3, GND, VFB, MPPSET, SRP e SRN. Para PH/SW, usar ponto de teste com cuidado e somente para osciloscopio.')
    bullet(d,'Inclua no esquema o sentido da corrente, pino 1 dos conectores, tensao nominal dos capacitores e rating de fusivel/TVS. Isso evita ambiguidade na montagem.')
    add_heading(d,'11.3 Layout e EMC',2)
    add_table(d,['Zona','Regra pratica','Motivo'],[
        ['BQ24650 - loop quente','CIN, QH, QL e GND muito proximos; area PH pequena','Reduz EMI e ringing.'],
        ['Shunt RSR','Trilhas Kelvin separadas de SRP/SRN ao CI','Evita erro de corrente por queda na trilha.'],
        ['VFB/MPPSET/TS','Longe de PH e indutor; retorno limpo','Evita ruido alterar tensao/carga.'],
        ['U3 TPS62163','CIN, L3 e COUT no layout recomendado [R4]','Estabilidade e eficiencia.'],
        ['Bateria/corrente','Trilhas largas, cobre e conectores conforme corrente/temperatura','Limita queda e aquecimento.'],
    ],[1.5,3.1,2.2])
    note(d,'Revisao de layout','Antes de enviar Gerbers, compare visualmente as duas regioes de buck com os exemplos de layout dos datasheets [R1] e [R4]. Em conversores chaveados, layout e parte do circuito.', 'warn')

    add_heading(d,'12. Sequencia de montagem, bring-up e validacao')
    add_table(d,['Ordem','Teste','Criterio de aprovacao'],[
        ['1','Inspecao sem bateria/painel','Polaridade, curto VIN-GND, valores e footprints conferidos.'],
        ['2','Alimentar PCB principal por fonte limitada','5 V e 3V3 corretos, consumo sem carga esperado, sem aquecimento.'],
        ['3','Testar LVD com fonte de bancada simulando bateria','Desliga em V_cut e religa em V_reconnect sem oscilacao.'],
        ['4','Testar BQ com fonte limitada, sem bateria real inicialmente','VIN_CHG, gates e corrente conforme [R1]; nenhuma sobretensao.'],
        ['5','Conectar VRLA com F3 e corrente limitada','Vbulk/Vrecharge e corrente dentro do perfil do fabricante.'],
        ['6','Testar painel em condicoes reais','VIN proximo a MPPSET; nao excede 28 V mesmo no frio.'],
        ['7','Teste de noite / carga real','Autonomia medida >= requisito com margem.'],
        ['8','Falhas controladas','Desconectar painel, AUX, bateria e cabo; verificar estado seguro.'],
    ],[.6,3.15,3.05])
    add_p(d,'Nao comece o primeiro teste com painel e bateria ligados ao mesmo tempo. Uma fonte de bancada com limitacao de corrente e o primeiro instrumento de seguranca. Registre tensoes, temperaturas e correntes em uma planilha de validacao.')

    d.add_page_break()
    add_heading(d,'13. Lista de componentes - ponto de partida, nao BOM final')
    add_table(d,['Bloco','Componente / referencia','Criterios para comprar','Status'],[
        ['Carregador','U1: BQ24650RVAR (TI)','16-QFN 3,5 x 3,5 mm; confirmar capacidade de montagem','Definido como arquitetura'],
        ['LVD','U2: TPS3700DDCT (TI)','SOT-23-6, 1,8-18 V, open-drain','Definido como arquitetura'],
        ['Buck 5 V','U3: TPS62163DSGR (TI)','5 V fixo, 1 A, WSON; montagem reflow','Definido; avaliar hand-solder'],
        ['Buck 5 V alternativa','TPS62160DGKR (TI)','VSSOP ajustavel, mais acessivel para solda; divisor para 5 V','Alternativa se encapsulamento pesar'],
        ['3V3 LDO','MCP1700-3302 (Microchip)','Baixo Iq; verificar corrente/potencia/condensadores','Condicional a I3V3'],
        ['OR entrada','Schottky ou PMOS ideal diode','Avaliar P = I x VF ou I2R e temperatura','Escolher apos Icharge'],
        ['Bateria','VRLA AGM 12 V, capacidade calculada','Perfil de carga e dimensoes validados','A selecionar'],
        ['Painel','12 V nominal','Vmp para MPPSET; Voc frio < limite do BQ','A selecionar'],
    ],[1.2,2.25,2.5,1.1])
    note(d,'Custo x confiabilidade','Para primeira placa, economizar em circuitos de baixa tensao e encapsulamentos pode ser razoavel; nao economize em fusivel da bateria, conector/cabo de potencia, distancias de isolamento, TVS e validação do perfil VRLA. Esses itens evitam as falhas mais energeticas.', 'risk')

    add_heading(d,'14. Referencias tecnicas oficiais')
    refs=[
        ('[R1] TI - BQ24650 datasheet e Typical Application','https://www.ti.com/document-viewer/BQ24650/datasheet'),
        ('[R2] TI SLVA437A - Using the bq24650 to Charge a Sealed, Lead-Acid Battery','https://www.ti.com/lit/an/slva437a/slva437a.pdf'),
        ('[R3] TI - TPS3700 datasheet, divisor, histerese e layout','https://www.ti.com/document-viewer/TPS3700/datasheet'),
        ('[R4] TI - TPS6216x datasheet, TPS62163 5 V, componentes e layout','https://www.ti.com/document-viewer/TPS62163/datasheet'),
        ('[R5] Microchip - MCP1700 datasheet','https://ww1.microchip.com/downloads/aemDocuments/documents/APID/ProductDocuments/DataSheets/MCP1700-Data-Sheet-20001826F.pdf'),
        ('[R6] TI - Ideal diode / ORing controllers overview','https://www.ti.com/product-category/power-management/power-protection-switches-controllers/ideal-diode-oring-controllers/overview.html'),
        ('[R7] TI - TPS3700EVM-114, referencia de avaliacao em SOT-23','https://www.ti.com/tool/TPS3700EVM-114'),
        ('[R8] TI - produto TPS62163 e codigo TPS62163DSGR','https://www.ti.com/product/TPS62163/part-details/TPS62163DSGR'),
    ]
    for title,url in refs:
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run(title+'\n'+url); r.font.name='Arial'; r.font.size=Pt(9.2); r.font.color.rgb=RGBColor.from_string(NAVY)
    add_p(d,'As referencias [R1]-[R4] sao a fonte de verdade para pinagem, valores iniciais, limites absolutos, rede de compensacao e layout. Esta referencia organiza a arquitetura do TCC, mas nao substitui a revisao desses documentos na versao vigente antes da compra/fabricacao.')

    add_heading(d,'15. Checklist antes de liberar a PCB')
    checks=[
        'Modelo de VRLA escolhido e perfil Vbulk/Vrecharge/corrente/temperatura copiado do fabricante.',
        'Painel escolhido: Vmp programado em MPPSET; Voc no frio e transitorios dentro do BQ24650.',
        'Consumo medio e picos medidos em prototipo; bateria e painel recalculados com margem.',
        'Fusivel F3 posicionado proximo ao terminal positivo da bateria; ratings de cabos/conectores verificados.',
        'P4 e J_PWR nao estao unidos diretamente: JP1, OR ou variante foi escolhido e documentado.',
        'TPS3700 LVD testado com fonte de bancada sob carga real; corte e religamento sem chatter.',
        'Layout de BQ24650 e TPS62163 comparado aos layouts de referencia; vias/planos/loops revisados.',
        'Test points, DNPs e procedimentos de bring-up incluidos no projeto KiCad e na BOM.',
    ]
    for c in checks: bullet(d,'[ ] '+c)
    d.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
